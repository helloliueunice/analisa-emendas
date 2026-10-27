"""Extração de texto de PDF, DOCX e TXT."""
from __future__ import annotations

import io


def extrair_texto(nome: str, dados: bytes) -> str:
    """Retorna o texto de um arquivo a partir do nome (extensão) e dos bytes."""
    nome_l = nome.lower()
    try:
        if nome_l.endswith(".pdf"):
            from pypdf import PdfReader

            leitor = PdfReader(io.BytesIO(dados))
            return "\n".join((p.extract_text() or "") for p in leitor.pages)
        if nome_l.endswith(".docx"):
            import docx

            doc = docx.Document(io.BytesIO(dados))
            partes = [p.text for p in doc.paragraphs]
            for tabela in doc.tables:
                for linha in tabela.rows:
                    partes.append(" | ".join(c.text for c in linha.cells))
            return "\n".join(partes)
        if nome_l.endswith((".txt", ".md", ".csv", ".json")):
            return dados.decode("utf-8", errors="ignore")
    except Exception as exc:  # arquivo corrompido, protegido etc.
        return f"[ERRO NA EXTRAÇÃO: {exc}]"
    return ""


def texto_parece_digitalizado(texto: str, n_paginas: int = 1) -> bool:
    """Heurística: pouco texto por página sugere PDF escaneado (precisa de OCR)."""
    return len(texto.strip()) < 80 * max(n_paginas, 1)
