"""Consolida os achados dos módulos em uma minuta de parecer técnico
(Markdown e DOCX). A minuta é um rascunho: a decisão é do analista."""
from __future__ import annotations

import io
from datetime import date

from .modelos import Achado, ResultadoModulo, Severidade

ORDEM = [Severidade.CRITICA, Severidade.ALTA, Severidade.MEDIA, Severidade.INFO]


def recomendacao_geral(resultados: list[ResultadoModulo], achados_ia: list[Achado] | None = None) -> str:
    todos = [a for r in resultados for a in r.achados] + [a for a in (achados_ia or []) if a.origem == "ia"]
    if any(a.severidade == Severidade.CRITICA and a.origem == "regra" for a in todos):
        return "NÃO APTO – há pendências impeditivas a sanar (diligência obrigatória)."
    if any(a.severidade in (Severidade.CRITICA, Severidade.ALTA) for a in todos):
        return "APTO COM RESSALVAS – recomenda-se diligência antes da aprovação."
    return "APTO – não foram identificadas pendências relevantes pelas verificações automatizadas."


def gerar_markdown(meta: dict, resultados: list[ResultadoModulo], achados_ia: list[Achado] | None = None,
                   sintese_ia: str = "") -> str:
    hoje = date.today().strftime("%d/%m/%Y")
    L = [f"# Minuta de Parecer Técnico – {meta.get('proposta', 'Proposta')}", "",
         f"**Proponente:** {meta.get('proponente', '—')}  ",
         f"**Emenda:** {meta.get('emenda', '—')}  ",
         f"**Objeto:** {meta.get('objeto', '—')}  ",
         f"**Data:** {hoje}  ", f"**Analista:** {meta.get('analista', '—')}", "",
         "## 1. Conclusão preliminar", "", f"**{recomendacao_geral(resultados, achados_ia)}**", ""]
    if sintese_ia:
        L += ["## 2. Síntese da análise (gerada por IA – revisar)", "", sintese_ia, ""]
    L += ["## 3. Resultado por módulo", ""]
    for r in resultados:
        L.append(f"- **{r.modulo}:** {len(r.achados)} apontamento(s); índice de risco {r.nota_risco}/100")
    L += ["", "## 4. Apontamentos", ""]
    todos = [a for r in resultados for a in r.achados] + list(achados_ia or [])
    for sev in ORDEM:
        grupo = [a for a in todos if a.severidade == sev]
        if not grupo:
            continue
        L += [f"### {sev.value} ({len(grupo)})", ""]
        for a in grupo:
            tag = " _(sugerido por IA)_" if a.origem == "ia" else ""
            L.append(f"- **[{a.modulo}] {a.regra}**{tag}: {a.mensagem}")
            if a.recomendacao:
                L.append(f"  - Providência: {a.recomendacao}")
            if a.fundamento:
                L.append(f"  - Fundamento: {a.fundamento}")
        L.append("")
    L += ["## 5. Ressalva", "",
          "Esta minuta foi gerada por ferramenta de apoio à análise. As verificações automatizadas não "
          "substituem o juízo técnico do analista nem a manifestação jurídica, e os fundamentos indicados "
          "devem ser conferidos com o normativo vigente.", ""]
    return "\n".join(L)


def gerar_docx(markdown: str) -> bytes:
    import docx

    d = docx.Document()
    for linha in markdown.splitlines():
        if linha.startswith("# "):
            d.add_heading(linha[2:], 0)
        elif linha.startswith("## "):
            d.add_heading(linha[3:], 1)
        elif linha.startswith("### "):
            d.add_heading(linha[4:], 2)
        elif linha.strip().startswith("- "):
            nivel = "List Bullet 2" if linha.startswith("  ") else "List Bullet"
            d.add_paragraph(linha.strip()[2:].replace("**", "").replace("_", ""), style=nivel)
        elif linha.strip():
            d.add_paragraph(linha.replace("**", "").rstrip())
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()
